import streamlit as st
from fpdf import FPDF
from fuzzywuzzy import fuzz  # fuzzywuzzy aik Python library hai jo do strings (matlab do lafzon ya jumlon) ko milakar unki milti-julti percentage batata hai
import re   # re text ke andar se pattern dhoondhne, unwanted characters hatane, ya validation karne mein madad deta hai
import docx
from pdfminer.high_level import extract_text  # ye line PDF document ke andar likha hua text extract karne mein help karti hai 
import os
import matplotlib.pyplot as plt  # data analyze
import spacy  # spacy text ko samajhne wali AI banata hai ye batata ha k Apple fruit ha ya company

nlp = spacy.load("en_core_web_sm")  # spacy ke pretrained Eng model (en_core_web_sm) load krta ha ta hum us model se text ko process kr skain.


# --- User Auth ---
def login(username, password):
    users = {"Madam": "madam4321", "hr_user": "tecrix_hr"}
    return users.get(username) == password

# --- File Extractors ---
def extract_text_from_pdf(file):
    return extract_text(file)

def extract_text_from_docx(file):
    doc = docx.Document(file)
    return "\n".join([para.text for para in doc.paragraphs])

# --- NLP Extractors ---
def extract_skills(text):
    skills_list = ['python', 'sql', 'excel', 'power bi', 'machine learning', 'communication', 'teamwork', 'deep learning', 'ai']
    return [s for s in skills_list if s.lower() in text.lower()]

def extract_education(text):
    edu_keywords = ['bachelor', 'master', 'phd', 'bs', 'ms', 'bsc', 'msc']
    return [e for e in edu_keywords if e in text.lower()]

def extract_experience(text):
    match = re.search(r'(\d+)\+?\s+years?', text, re.I)
    return match.group(0) if match else "Not mentioned"


def extract_phone(text):
    phone_regex = re.findall(r'(?:(?:\+92|0092|0)\s?\d{3}[-\s]?\d{7})|(?:\d{11})|(?:\d{4}[-\s]?\d{7})', text)
    return phone_regex[0] if phone_regex else "Not found"


def extract_name(text):
    doc = nlp(text)
    for ent in doc.ents:
        if ent.label_ == "PERSON":
            return ent.text
    return "Unknown"

def extract_entities(text):
    doc = nlp(text)
    return {ent.label_: ent.text for ent in doc.ents if ent.label_ in ['ORG', 'DATE', 'PERSON']}

# --- Scorer & Role Suggestion ---
def match_score(resume, jd):
    score = fuzz.token_set_ratio(resume.lower(), jd.lower())
    if score >= 85:
        return score, "🔥 Strong Match"
    elif score >= 60:
        return score, "🧐 Needs Review"
    else:
        return score, "❌ Not Suitable"

def suggest_roles(skills):
    if 'python' in skills and 'sql' in skills:
        return ["Data Analyst", "BI Developer", "ML Engineer"]
    elif 'excel' in skills:
        return ["Operations Assistant", "Data Entry", "Project Coordinator"]
    else:
        return ["General Role Suggestion: Admin, HR Assistant"]

# --- PDF Report ---
def generate_report(name,phone, skills, education, experience, score, recommendation, missing):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt=f"Resume Report for {name}", ln=True, align='C')
    pdf.ln(10)
    pdf.cell(200, 10, txt=f"Skills: {', '.join(skills)}", ln=True)
    pdf.cell(200, 10, txt=f"Phone: {phone}", ln=True)
    pdf.cell(200, 10, txt=f"Education: {', '.join(education)}", ln=True)
    pdf.cell(200, 10, txt=f"Experience: {experience}", ln=True)
    pdf.cell(200, 10, txt=f"Missing Skills: {', '.join(missing)}", ln=True)
    pdf.cell(200, 10, txt=f"Match Score: {score}%", ln=True)
    clean_rec = recommendation.encode("ascii", "ignore").decode()
    pdf.cell(200, 10, txt=f"Recommendation: {clean_rec}", ln=True)
    os.makedirs("reports", exist_ok=True)
    safe_name = re.sub(r'[\\/*?:"<>|\n\r]', "", name.strip())
    file_path = f"reports/{safe_name}_report.pdf"
    pdf.output(file_path)
    return file_path

# --- Streamlit App ---
st.set_page_config("HR Resume Analyzer", page_icon="📄")
st.title("🔐 HR Resume Analyzer Login")

if "logged_in" not in st.session_state:
    st.session_state.logged_in = False

if not st.session_state.logged_in:
    username = st.text_input("Username")
    password = st.text_input("Password", type="password")
    if st.button("Login"):
        if login(username, password):
            st.session_state.logged_in = True
            st.success("✅ Login Successful!")
        else:
            st.error("❌ Invalid username or password")

if st.session_state.logged_in:
    st.title("📄 Resume Analyzer")
    resume_file = st.file_uploader("Upload Resume (PDF/DOCX)", type=["pdf", "docx"])
    job_desc = st.text_area("Paste Job Description")

    disabled_btn = not (resume_file and job_desc)
    if st.button("🔍 Start Analysis", disabled=disabled_btn):


            resume_text = extract_text_from_pdf(resume_file) if resume_file.name.endswith("pdf") else extract_text_from_docx(resume_file)
            skills = extract_skills(resume_text)
            education = extract_education(resume_text)
            experience = extract_experience(resume_text)
            phone = extract_phone(resume_text)
            name = extract_name(resume_text)
            entities = extract_entities(resume_text)

            jd_skills = extract_skills(job_desc)
            missing = [s for s in jd_skills if s not in skills]
            match_percent = int((len(jd_skills) - len(missing)) / len(jd_skills) * 100) if jd_skills else 0
            score, rec = match_score(resume_text, job_desc)

            st.markdown("🧑‍💼 Resume Info")
            st.markdown(f"**Name:** {name}")
            # st.markdown(f"**Phone:** {phone}")

            st.markdown("🧠 Skills")
            for skill in skills:
                st.markdown(f"- {skill.capitalize()}")



            st.markdown("🎓 Education")
            for edu in education:
                st.markdown(f"- {edu.upper()}")

            st.markdown("💼 Experience")
            st.markdown(f"- {experience}")

            st.markdown("🏢 Entities (ORG/DATE)")
            for k, v in entities.items():
                st.markdown(f"- **{k}**: {v}")

            st.markdown("---")
            st.markdown("📊 JD Matching")
            st.markdown(" Required Skills from JD")
            for js in jd_skills:
                st.markdown(f"- {js}")

            st.markdown("Missing Skills")
            if missing:
                for ms in missing:
                    st.markdown(f"- {ms}")
            else:
                st.markdown("No missing skills! You're a perfect fit. 💯")

            st.markdown(f"📈 Skill Match %: `{match_percent}%`")
            st.markdown(f"🧪 Fit Score: `{score}% – {rec}`")

            st.subheader("👾 Suggested Job Roles")
            st.write(suggest_roles(skills))

            st.markdown("---")
            st.subheader("📅 Download PDF Report")
            report_path = generate_report(name, phone, skills, education, experience, score, rec, missing)
            with open(report_path, "rb") as f:
                st.download_button(" Download Report", f, file_name=os.path.basename(report_path))
         

